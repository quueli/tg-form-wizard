import abc
from datetime import datetime
from typing import Iterable

from data import questionnaire as q

# (answer key, labels to look the value up in, or None for raw text)
COLUMNS: list[tuple[str, dict | None]] = [
    ("segment", q.SEGMENTS),
    ("category", q.CATEGORIES),
    ("group", q.GROUPS),
    ("variant", q.VARIANTS),
    ("item_a_include", q.ITEMS["a"]),
    ("item_a_exclude", q.ITEMS["a"]),
    ("item_b_include", q.ITEMS["b"]),
    ("item_c_include", q.ITEMS["c"]),
    ("size", q.SIZES),
    ("level", q.LEVELS),
    ("speed", q.SPEEDS),
    ("tags", q.TAGS),
    ("priority", q.PRIORITIES),
    ("track", q.TRACKS),
    ("option_final", q.OPTIONS_FINAL),
    ("flags", q.FLAGS),
    ("customer_type", q.CUSTOMER_TYPES),
    ("notes", None),
    ("contact", None),
]

HEADER = [key.replace("_", " ") for key, _ in COLUMNS]


def _label(value, lookup: dict | None) -> str:
    if value is None or value == []:
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(_label(v, lookup) for v in value)
    if lookup is not None:
        return lookup.get(value, str(value))
    return str(value)


def answers_to_row(answers: dict) -> list[str]:
    return [_label(answers.get(key), lookup) for key, lookup in COLUMNS]


class SheetWriter(abc.ABC):
    @abc.abstractmethod
    def append(self, row: Iterable[str]) -> None: ...

    def append_answers(self, answers: dict) -> None:
        self.append([datetime.now().isoformat(timespec="seconds"), *answers_to_row(answers)])


class XlsxSheetWriter(SheetWriter):
    def __init__(self, path: str, sheet_title: str = "responses"):
        from openpyxl import Workbook, load_workbook

        self.path = path
        try:
            self._wb = load_workbook(path)
            self._ws = self._wb.active
        except FileNotFoundError:
            self._wb = Workbook()
            self._ws = self._wb.active
            self._ws.title = sheet_title
            self._ws.append(["timestamp", *HEADER])

    def append(self, row: Iterable[str]) -> None:
        self._ws.append(list(row))
        self._wb.save(self.path)


class GoogleSheetsWriter(SheetWriter):
    def __init__(self, spreadsheet_id: str, credentials_file: str, worksheet: str = "responses"):
        import gspread  # only dep of this writer, keep it out of the xlsx path

        client = gspread.service_account(filename=credentials_file)
        book = client.open_by_key(spreadsheet_id)
        try:
            self._ws = book.worksheet(worksheet)
        except gspread.WorksheetNotFound:
            self._ws = book.add_worksheet(worksheet, rows=1, cols=len(HEADER) + 1)
            self._ws.append_row(["timestamp", *HEADER])

    def append(self, row: Iterable[str]) -> None:
        self._ws.append_row(list(row), value_input_option="USER_ENTERED")
