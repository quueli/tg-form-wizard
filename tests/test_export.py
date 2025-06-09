from services.export import HEADER, XlsxSheetWriter, answers_to_row


def test_answers_map_to_labels():
    answers = {
        "segment": "segment_1",
        "category": "category_1",
        "item_a_include": ["a1", "a2"],
        "notes": "free text",
    }
    row = dict(zip(HEADER, answers_to_row(answers)))
    assert row["segment"] == "Segment 1"
    assert row["category"] == "Category 1"
    assert row["item a include"] == "Option A1, Option A2"
    assert row["notes"] == "free text"


def test_missing_answers_become_empty():
    row = answers_to_row({})
    assert all(cell == "" for cell in row)


def test_xlsx_writer_writes_header(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "out.xlsx"
    writer = XlsxSheetWriter(str(path))
    writer.append_answers({"segment": "segment_1"})

    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "timestamp"
    assert ws.cell(row=2, column=2).value == "Segment 1"


def test_xlsx_writer_appends_to_existing(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "out.xlsx"
    XlsxSheetWriter(str(path)).append_answers({"segment": "segment_1"})
    XlsxSheetWriter(str(path)).append_answers({"segment": "segment_1"})

    ws = load_workbook(path).active
    assert ws.max_row == 3
